"""Erzeugt vorlagen/ZSP-Manuskriptvorlage.docx (Hausformat: Arial 11 pt, Zeilenabstand 1,5, linksbündig, Seitenzahl in der Fußzeile)."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent.parent / "vorlagen" / "ZSP-Manuskriptvorlage.docx"
GRAU = RGBColor(0x4B, 0x56, 0x61)


def arial(run, size=11, bold=None, italic=None, color=None):
    run.font.name = "Arial"; run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color is not None: run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rPr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), "Arial")


doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(2.5); sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.0)

for name in ("Normal", "Heading 1", "Heading 2", "Title"):
    st = doc.styles[name]
    st.font.name = "Arial"; st.font.size = Pt(11); st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(6)
doc.styles["Heading 1"].font.bold = True
doc.styles["Heading 2"].font.bold = True; doc.styles["Heading 2"].font.italic = True


def para(text="", bold=False, italic=False, color=None, style=None, after=6):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(after)
    if text:
        arial(p.add_run(text), bold=bold, italic=italic, color=color)
    return p


def hinweis(text):
    return para(text, italic=True, color=GRAU)


def heading(text, level=1):
    p = doc.add_heading(level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    arial(p.add_run(text), bold=True, italic=(level == 2), color=RGBColor(0, 0, 0))
    return p


# Fußzeile mit Seitenzahl (PAGE-Feld)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run(); arial(r)
for tag, attr in (("w:fldChar", {"w:fldCharType": "begin"}), ("w:instrText", None), ("w:fldChar", {"w:fldCharType": "end"})):
    el = OxmlElement(tag)
    if attr:
        for k, v in attr.items(): el.set(qn(k), v)
    else:
        el.set(qn("xml:space"), "preserve"); el.text = "PAGE"
    r._element.append(el)
hp = sec.header.paragraphs[0]
arial(hp.add_run("Zeitschrift für Sinnzentrierte Psychologie · Manuskript zur Begutachtung"), size=11, color=GRAU)

# Titelblatt
para("Zeitschrift für Sinnzentrierte Psychologie", bold=True)
para("Manuskriptvorlage", color=GRAU, after=18)
para("Titel des Beitrags", bold=True)
hinweis("Prägnanter Titel, höchstens 15 Wörter. Ein Untertitel ist möglich.")
para("Untertitel", bold=True)
para("Kurztitel für die Kopfzeile", bold=True)
hinweis("Höchstens 60 Zeichen einschließlich Leerzeichen.")
para("Rubrik", bold=True)
hinweis("Übersichtsarbeit, Originalarbeit, Theoretischer Beitrag, Methodenbeitrag, Praxis und Lehre oder Rezension.")
para("Wortzahl", bold=True)
hinweis("Gesamtzahl der Wörter einschließlich Literatur, ohne Abstracts.")
para("Hinweis zur Verblindung", bold=True)
hinweis("Namen, Institutionen und Danksagungen stehen ausschließlich im separaten Titelblatt, das Sie mit der Einreichung senden. Dieses Manuskript enthält keine Angaben, die auf die Autorinnen und Autoren schließen lassen. Eigene Arbeiten zitieren Sie in der dritten Person.")

doc.add_page_break()
heading("Zusammenfassung")
hinweis("Strukturiert, höchstens 250 Wörter: Hintergrund, Fragestellung, Methode, Ergebnisse, Schlussfolgerungen. Bei theoretischen Beiträgen: Hintergrund, Argumentation, Schlussfolgerungen.")
para("Schlüsselwörter:", bold=True)
hinweis("Fünf Begriffe, durch Kommata getrennt.")
heading("Abstract")
hinweis("Englische Fassung der Zusammenfassung, höchstens 250 Wörter.")
para("Keywords:", bold=True)

doc.add_page_break()
heading("1 Einleitung")
hinweis("Problemstellung, Relevanz und Ziel des Beitrags. Zitieren Sie nach APA 7, zum Beispiel (Frankl, 1946) oder Steger et al. (2006).")
heading("2 Theoretischer Hintergrund")
hinweis("Stand der Forschung und begriffliche Klärungen. Unterabschnitte werden zweistufig nummeriert.")
heading("2.1 Unterabschnitt", level=2)
heading("3 Methode")
hinweis("Nur für empirische Arbeiten: Design, Stichprobe, Instrumente, Ablauf, Auswertung, Ethikvotum, gegebenenfalls Präregistrierung.")
heading("4 Ergebnisse")
hinweis("Tabellen und Abbildungen werden fortlaufend nummeriert und im Text erwähnt. Effektstärken und Konfidenzintervalle werden berichtet.")
heading("5 Diskussion")
hinweis("Einordnung der Befunde, Grenzen des Beitrags, Folgerungen für Forschung und Praxis.")
heading("6 Fazit")

doc.add_page_break()
heading("Erklärungen")
for t, h in [
    ("Finanzierung", "Angabe aller Förderungen oder die Erklärung, dass keine Förderung vorlag."),
    ("Interessenkonflikte", "Offenlegung finanzieller und nicht finanzieller Interessen oder die Erklärung, dass keine bestehen."),
    ("Ethik", "Bei Forschung mit Menschen: zuständige Ethikkommission und Votum, informierte Einwilligung."),
    ("Verfügbarkeit von Daten und Analysecode", "Ort der Daten und des Codes oder Begründung, warum sie nicht geteilt werden können."),
    ("Beiträge der Autorinnen und Autoren", "Nach der CRediT-Taxonomie."),
    ("Einsatz generativer KI", "Art und Umfang des Einsatzes, zum Beispiel für sprachliche Überarbeitung, oder die Erklärung, dass keine KI eingesetzt wurde. KI-Werkzeuge können keine Autorschaft beanspruchen."),
]:
    para(t, bold=True, after=0)
    hinweis(h)

heading("Literatur")
hinweis("Alphabetisch nach APA 7, mit DOI, sofern vorhanden. Beispiele:")
for ref in [
    "Frankl, V. E. (1946). Ärztliche Seelsorge. Grundlagen der Logotherapie und Existenzanalyse. Deuticke.",
    "Steger, M. F., Frazier, P., Oishi, S. & Kaler, M. (2006). The Meaning in Life Questionnaire: Assessing the presence of and search for meaning in life. Journal of Counseling Psychology, 53(1), 80–93.",
    "Vos, J., Craig, M. & Cooper, M. (2015). Existential therapies: A meta-analysis of their effects on psychological outcomes. Journal of Consulting and Clinical Psychology, 83(1), 115–128.",
]:
    p = para(ref)
    p.paragraph_format.left_indent = Cm(1.25); p.paragraph_format.first_line_indent = Cm(-1.25)

# Hausformat absichern
for p in doc.paragraphs:
    p.paragraph_format.line_spacing = 1.5
    if p.alignment != WD_ALIGN_PARAGRAPH.CENTER:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:
        arial(r, size=11, bold=r.font.bold, italic=r.font.italic, color=r.font.color.rgb if r.font.color and r.font.color.type else None)

doc.core_properties.title = "Manuskriptvorlage · Zeitschrift für Sinnzentrierte Psychologie"
doc.core_properties.author = "Institut für Sinnzentrierte Psychologie"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)

# Prüfung
chk = Document(OUT)
bad = [(i, r.text[:30]) for i, p in enumerate(chk.paragraphs) for r in p.runs if r.font.name != "Arial" or r.font.size != Pt(11)]
sp = {p.paragraph_format.line_spacing for p in chk.paragraphs}
print(OUT.name, "Absätze:", len(chk.paragraphs), "Abweichungen:", bad, "Zeilenabstände:", sp)
